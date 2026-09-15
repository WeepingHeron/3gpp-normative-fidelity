"""Preserve v4; publish the sealed v2 reevaluation in the existing KIPS layout."""
from copy import deepcopy
from pathlib import Path
import hashlib, json
from docx import Document
from docx.shared import Inches, Pt
from revise_paper_template_20260906 import replace_text
ROOT=Path(__file__).resolve().parents[1]
BASE=ROOT/'outputs/paper'
SRC=BASE/'ACK2026_3GPP_normative_fidelity_draft_v4.docx'
OUT=BASE/'ACK2026_3GPP_normative_fidelity_validated_v5.docx'
EVIDENCE=ROOT/'outputs/review/v2_cohort_analysis_20260909'

def main():
    assert hashlib.sha256(SRC.read_bytes()).hexdigest()=='cf1a36a8a4deff6af047539d9c4e954c12d24011cce56011a86ad02a77dd6dc8'
    records=[json.loads(x) for x in (EVIDENCE/'condition_results.jsonl').read_text().splitlines()]
    primary=[x for x in records if x['cohort']=='PRIMARY_STRICT_49']
    chart=BASE/"cohort_bounds_paper_v5.png"; assert chart.exists()
    d=Document(SRC)
    edits={
      4:'연구 질문은 원문 단독(C1)에 정답 구조(C2) 또는 자동 추출 구조(C3)를 추가할 때 한국어 설명의 완전 의미 보존율이 어떻게 달라지는가이다. 정답 규범 유형·부정만 추가한 C1M도 비교한다. 본고는 동결된 응답의 표본 적격성 재점검과 지연 재평가를 보고하며, 실제 검색 정확도나 3GPP 전체 일반화는 다루지 않는다.',
      6:'원문은 ETSI 공식 배포본인 TS 23.501 및 TS 24.501 V18.12.0 Release 18로 고정했다[8], [9]. 최초 표본은 7,727개 후보에서 선정한 72개 문장(규범 유형 6종별 12개, 44개 조항)이다. 생성 후 적격성을 다시 감사해 일반 적격 50건, 조건부 예외 1건, 부적격 10건, 보류 11건으로 구분했다. 주 분석은 일반 적격 중 C2 정답 입력 오염 1건(MAY-001)을 제외한 49건·33개 조항이다. 제외 기준은 사후에 마련해 이번 재평가 전에 동결했다. 원문·질문·기존 응답은 수정하거나 재생성하지 않았다.',
      7:'C0는 원문 없는 질문 조건이다. C1은 원문과 의미 보존 지시를, C2는 원문과 정답 actor·modal·polarity·condition·action·exception을 받는 oracle structured 조건이다. C3는 질문·정답 없이 원문에서 여섯 필드를 자동 추출한 뒤 설명하고 불명확하면 보류한다. C1M은 C1에 정답 modal·polarity만 추가한다. 모델은 gpt-5.4-mini-2026-03-17, temperature 0, reasoning effort none, 출력 상한 600 tokens이며 C3 추출은 1,200 tokens다. C1/C2는 9월 3일, C3/C1M은 9월 6일 생성했다. C2와 C1M은 정답 힌트를 쓰므로 자동 추출 성능과 구분한다.',
      8:'행위자·규범 강도와 부정·조건·행동·예외를 모두 보존하고 미지원 규범을 추가하지 않으면 P(성공), 명확한 위반이면 F(실패), 의미 판정을 확정할 수 없으면 U(미확정)로 평가했다. 미생성 NG는 의미 오류와 구분하되 분모에 포함해 성공 0건으로 계산한다. 51건×5조건의 255개 슬롯을 동일 Codex가 9월 8일 평가하고 24시간 이상 후 재검토해 불일치 5건을 조정했다. 기존 결과를 아는 동일 AI의 검토이므로 독립 맹검이나 전문가 검증은 아니다. 확정 판정을 고정하고 U를 모두 실패 또는 성공으로 둘 때의 P/N~(P+U)/N만 보고한다. 이는 신뢰구간이 아니며 U가 남은 추론 통계는 보류했다.',
      10:'<표 1> 주 분석 49건의 완전 의미 보존 결과',
      11:'표 1의 범위는 미확정 판정에 따른 논리적 상·하한이다. C1은 P 38건·U 1건, C2는 P 44건·U 2건이며 C2−C1은 +10.20~+16.33%p다. C3−C1은 +4.08~+10.20%p, C1M−C1은 +8.16~+14.29%p다. 주 분석의 U 7건과 C3 미생성 1건은 모두 분모에 남겼다. C0의 0/49는 원문 없는 조건에서 이 완전 보존 기준을 충족하지 못한 결과다.',
      13:'(그림 1) 표본별 C2−C1 논리적 범위(신뢰구간 아님)',
      14:'민감도 분석은 주 분석에 입력 오염 사례를 복원한 Gold 50, 조건부 예외(SDN-011)를 복원한 Exception 50, 둘 다 복원한 Both 51이다(그림 1). 평가 정답의 보정은 원래 C2 입력의 오염을 제거하지 않는다. 네 표본 모두 C2·C3·C1M의 C1 대비 차이 하한은 양수지만, 이는 확정 판정을 고정한 기술적 비교이며 모집단의 우월성을 입증하지 않는다. 표본이 서로 중첩되므로 독립 반복 실험도 아니다.',
      15:'C2−C1M은 주 분석에서 −2.04~+6.12%p이고 세 민감도 표본에서도 0을 포함한다. 따라서 전체 정답 구조와 유형·부정 힌트의 순위 또는 동등성을 확정할 수 없다. 출처 참조만으로 판단이 해소되지 않은 사례는 U로 유지했다. JSON의 규범 유형이나 인용 ID가 맞아도 설명의 주체·행동·조건 연결을 대신 검증할 수 없으므로, 메타데이터 일치를 의미 보존 성공으로 간주하지 않았다.',
      17:'동결 응답의 사후 재평가에서 정답 구조·자동 추출 구조·규범 유형 힌트 조건은 원문 단독보다 높은 완전 보존 범위를 보였다. 다만 이 결과는 확정 판정과 선정 표본에 조건부인 기술적 관찰이다. 규범 유형 이름뿐 아니라 그것이 적용되는 주체·행동과 조건·예외의 범위를 함께 평가해야 하며, 단일 AI의 재검토 일치는 정답 타당도를 보장하지 않는다.',
      18:'본 연구는 단일 모델·두 규격의 고정 버전과 Codex 기반 정답·평가에 한정된다. 사후 표본 제외, 조항 내 의존성, 미확정 판정이 남아 통계적 유의성이나 일반화를 주장하지 않는다. C2와 C1M은 지시문도 달라 순수한 구조 효과를 분리하지 못하며 생성일 차이도 남는다. 후속 개발 사례는 본 분석에 합치지 않았다. 독립 전문가 판정과 개발에 사용하지 않은 새 조항 검증이 필요하다.'
    }
    assert len(d.paragraphs)==29
    for i,text in edits.items():replace_text(d.paragraphs[i],text)
    abstract=d.tables[0].cell(8,0).paragraphs[1]
    replace_text(abstract,'3GPP 영문 규격의 한국어 설명에서 규범 강도·행위자·조건·행동·예외의 보존을 평가했다. TS 23.501과 TS 24.501 V18.12.0의 동결 응답을 사후 적격성 감사와 동일 AI의 지연 재검토로 재평가했다. 주 분석 49건에서 완전 보존 성공은 원문 단독(C1) 38건, 정답 구조(C2) 44건, 자동 추출 구조(C3) 41건, 정답 유형·부정 힌트(C1M) 43건이었다. 미확정을 유지한 보존율 범위는 각각 77.55~79.59%, 89.80~93.88%, 83.67~87.76%, 87.76~91.84%다. 확정 판정을 고정하면 세 조건의 C1 대비 차이는 양수이나, 이 범위는 신뢰구간이 아니며 통계적 우월성이나 일반화를 뜻하지 않는다. 사후 표본 선정과 단일 AI 평가의 한계가 있어 독립 전문가 및 새 조항 검증이 필요하다.')
    table=d.tables[1];table._tbl.append(deepcopy(table.rows[-1]._tr))
    values=[['조건','P','F','U / NG','보존율 범위(%)']]
    for x in primary:
        values.append([x['condition'],str(x['PASS']),str(x['FAIL']),f"{x['UNRESOLVED']} / {x['NOT_SCORED']}",f"{100*x['lower_rate']:.2f}~{100*x['upper_rate']:.2f}"])
    widths=[.40,.32,.32,.59,1.57]
    table.autofit=False
    for col,w in zip(table.columns,widths):col.width=Inches(w)
    for row,vals in zip(table.rows,values):
        for cell,value,w in zip(row.cells,vals,widths):
            cell.width=Inches(w);replace_text(cell.paragraphs[0],value)
            for p in cell.paragraphs:
                for run in p.runs:run.font.size=Pt(9)
    p=d.paragraphs[12];p.clear();shape=p.add_run().add_picture(str(chart),width=Inches(3.2))
    shape._inline.docPr.set('descr','Four overlapping cohorts: logical C2 minus C1 bounds, not confidence intervals.')
    d.save(OUT)
    lines=['# '+d.tables[0].cell(0,0).text,'',d.tables[0].cell(4,0).text,'',d.tables[0].cell(2,0).text,'',d.tables[0].cell(6,0).text,'','## 요약','',abstract.text,'']
    for p in d.paragraphs:
        if not p.text.strip():continue
        lines += [('## ' if p.style.name=='논문 내용 제목' else '')+p.text,'']
        if p.text.startswith('<표 1>'):
            for i,row in enumerate(table.rows):
                lines.append('| '+' | '.join(c.text.replace('\n',' ') for c in row.cells)+' |')
                if i==0:lines.append('|---|---:|---:|---:|---:|')
            lines.append('')
        if p.text.startswith('(그림 1)'):lines+=['![C2−C1 논리적 범위](cohort_bounds_paper_v5.png)','']
    OUT.with_suffix('.md').write_text('\n'.join(lines))
    print(OUT)

if __name__=='__main__':main()
