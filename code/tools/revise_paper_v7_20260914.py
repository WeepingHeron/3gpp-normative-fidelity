from pathlib import Path
from docx import Document
from revise_paper_template_20260906 import replace_text
R=Path(__file__).resolve().parents[1];B=R/'outputs/paper';d=Document(B/'ACK2026_3GPP_normative_fidelity_validated_v6.docx');out=B/'ACK2026_3GPP_normative_fidelity_validated_v7.docx'
for p in d.paragraphs:
 if p.text.startswith('C0는'):
  replace_text(p,p.text.replace('C3는 질문·정답 없이 원문에서 여섯 필드를 자동 추출한 뒤 설명하고 불명확하면 보류한다.','C3의 추출은 원문만 사용한다. 설명 생성은 질문·원문·자동 추출 여섯 필드를 함께 사용하며 추출이 불명확하면 보류한다.'))
 if p.text.startswith('행위자·규범'):
  replace_text(p,p.text.replace('행위자·규범 강도와 부정·조건·행동·예외를 모두 보존하고','주 분석은 동결한 기준으로 설명의 행위자·규범 강도와 부정·조건·행동·예외를 보존하고'))
 if p.text.startswith('개발 원문 SMF'):
  replace_text(p,'개발 원문 SMF(§4.2.3)·AMF(§5.4.11.4)[8]·NSSAI(§4.6.2.2 및 Annex C)[9]에 문맥을 제공해 각 3회 직접 설명을 계획했다(9월 14일, 동일 모델, temperature 0, 2~4문장, 900토큰). 통신 미상 1개는 재시도하지 않았다. 추가 평가는 동결한 SMF 2개·NSSAI 4개·AMF 3개의 요구사항을 대상으로 한다. 선택적 능력 보유, SUPI 결합, 긴급·eCall의 비저장/임시 저장/삭제 사건, 위치 판단의 권고·위성 범위·망 설정을 확인한다. 복합 항목은 모든 요소를 보존해야 하며, 설명만 채점하고 메타데이터로 보충하지 않는다. 문맥 전체의 모든 명제를 망라한 지표는 아니다. 완료 8개는 같은 날 동일 AI의 재검토 후에도 완수 3·미완수 5였다(SMF 2/2, NSSAI 1/3, AMF 0/3).')
 if p.text.startswith('AMF 3개'):
  replace_text(p,'AMF 3개에서 강도 변경과 지정된 망 설정 조건의 미보존을 확인했다. 첫 응답의 “등록해제해서는 안 된다”는 should not의 권고 한정이 없다. NSSAI 2개는 eCall 예외 일부를 누락했다. 마지막 응답의 eCall 임시 저장은 메타데이터에만 남았고, 긴급 등록 절의 임시 저장을 eCall로 자동 확장해 읽지 않았다. 반면 “저장해야 한다는 권고”는 전체 표현이 권고를 명시해 강도 오류로 세지 않았다.')
 if p.text.startswith('본 연구는 단일'):
  replace_text(p,'단일 모델·고정 규격·동일 AI의 정답과 평가, 사후 표본 처리·조항 의존성·미확정 판정에 한정된다. 지시문과 생성일 차이로 순수한 구조 효과를 분리하지 못한다. 추가 진단은 알려진 원문의 반복이며 주 분석과 합산하지 않았다. 추가 재검토는 기존 판정을 아는 같은 AI가 수행했고 독립·지연 검증이 아니다. 통신 미상은 의미 오류와 구분했다. 통계적 일반화보다 독립 전문가와 새 조항 검증이 필요하다.')
for p in d.paragraphs:
 if p.text.startswith('정답 구조·자동 추출 구조·유형 힌트 조건의'):
  replace_text(p,'확정 판정 아래 세 조건의 C1 대비 차이는 양수지만 C2와 C1M의 우열은 정해지지 않았다. 규범 강도·검사 대상·조건·예외 오류는 유형 및 출처 필드와 별도로 자연어 설명의 명제 연결을 평가해야 함을 보여준다.')
 if p.text.startswith('민감도 분석은'):
  replace_text(p,'주 분석에 MAY-001을 더한 Gold 50, 조건부 예외 SDN-011을 더한 Exception 50, 둘 다 더한 Both 51을 비교했다(그림 1). 채점 정답 보정은 기존 C2 입력 오염을 제거하지 않는다. 네 표본에서 C2·C3·C1M의 C1 대비 차이 하한은 양수다. 이는 확정 판정에 조건부인 기술적 비교이며 모집단 우월성은 아니다. 중첩 표본을 독립 반복으로 보지 않는다.')
a=d.tables[0].cell(8,0).paragraphs[1];replace_text(a,a.text.replace('8개 설명의 1차 평가로','8개 설명의 지정 요구사항 재검토로'))
for p in d.paragraphs:
 if p.text.startswith('주 분석은 동결한 기준'):
  replace_text(p,'주 분석은 동결 기준에 따라 설명의 규범 요소를 모두 보존하고 미지원 규범 추가가 없으면 P, 명확한 위반은 F, 판정 미확정은 U다. 미생성 NG는 의미 오류와 구분하되 분모에 남겨 성공 0으로 센다. 합집합 51건×5조건을 동일 Codex가 9월 8일 평가하고 24시간 이상 후 재검토해 불일치 5건을 조정했다. 독립 맹검은 아니다. 확정 판정을 고정한 P/N~(P+U)/N은 U의 실패/성공 처리에 따른 논리적 범위이며 신뢰구간이 아니다. U가 남은 추론 통계는 보류했다.')
 if p.text.startswith('기존 연구는'):
  replace_text(p,'기존 연구는 3GPP 데이터셋·질의응답[2], 통신 규격 RAG[3], 한국어 통신 에이전트[4], 교차 참조·버전 검색[5], 규범·보안 추론[6]을 다뤘다. 법률 번역에는 계층적 제약으로 규범 강도를 보존하는 접근이 있다[7]. 본 연구는 고정 문장의 한국어 설명에서 규범 요소의 보존을 비교한다.')
d.save(out)
lines=['# '+d.tables[0].cell(0,0).text,'',d.tables[0].cell(4,0).text,'',d.tables[0].cell(2,0).text,'',d.tables[0].cell(6,0).text,'','## 요약','',a.text,'']
for p in d.paragraphs:
 if not p.text.strip():continue
 lines += [('## ' if p.style.name=='논문 내용 제목' else '')+p.text,'']
 if p.text.startswith('<표 1>'):
  for i,row in enumerate(d.tables[1].rows):
   lines.append('| '+' | '.join(c.text for c in row.cells)+' |')
   if i==0:lines.append('|---|---:|---:|---:|---:|')
  lines.append('')
 if p.text.startswith('(그림 1)'):lines+=['![논리적 범위](cohort_bounds_paper_v5.png)','']
out.with_suffix('.md').write_text('\n'.join(lines));print(out)
