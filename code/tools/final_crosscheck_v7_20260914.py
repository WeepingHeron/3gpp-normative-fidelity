from pathlib import Path
from collections import Counter
import hashlib,json,re,datetime
from docx import Document
import pdfplumber
R=Path(__file__).resolve().parents[1];B=R/'outputs/paper';O=R/'outputs/final_crosscheck_v7_20260914';stem='ACK2026_3GPP_normative_fidelity_validated_v7'
def h(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def rows(p):return list(map(json.loads,p.read_text().splitlines()))
md=(B/(stem+'.md')).read_text();d=Document(B/(stem+'.docx'));checks=[]
def check(name,ok):
 assert ok,name
 checks.append(name)
for p in d.paragraphs:
 if p.text.strip():check('DOCX paragraph present in Markdown: '+p.text[:16],p.text in md)
check('abstract present in Markdown',d.tables[0].cell(8,0).paragraphs[1].text in md)
slots=rows(R/'outputs/review/v2_cohort_analysis_20260909/bound_slots_255.jsonl');cr=rows(R/'outputs/review/v2_cohort_analysis_20260909/condition_results.jsonl')
for x in cr:
 rr=[r for r in slots if x['cohort']in r['cohorts'] and r['condition']==x['condition']];c=Counter(r['semantic_judgment']for r in rr)
 check('source counts '+x['cohort']+'/'+x['condition'],len(rr)==x['n'] and all(c[k]==x[k]for k in ['PASS','FAIL','UNRESOLVED','NOT_SCORED']))
prim=[x for x in cr if x['cohort']=='PRIMARY_STRICT_49']
expected=[['조건','P','F','U / NG','보존율 범위(%)']]+[[x['condition'],str(x['PASS']),str(x['FAIL']),f"{x['UNRESOLVED']} / {x['NOT_SCORED']}",f"{100*x['lower_rate']:.2f}~{100*x['upper_rate']:.2f}"]for x in prim]
check('DOCX table matches source counts',[[c.text for c in row.cells]for row in d.tables[1].rows]==expected)
for x in prim[1:]:check('abstract range '+x['condition'],f"{100*x['lower_rate']:.2f}~{100*x['upper_rate']:.2f}%"in d.tables[0].cell(8,0).text)
case={x['condition']:x for x in prim}
for a,b in [('C1','C2'),('C1','C3'),('C1','C1M'),('C1M','C2')]:
 lo=100*(case[b]['PASS']-case[a]['PASS']-case[a]['UNRESOLVED'])/49;hi=100*(case[b]['PASS']+case[b]['UNRESOLVED']-case[a]['PASS'])/49
 text=f'{lo:+.2f}~{hi:+.2f}%p'.replace('-','−');check('paper paired range '+b+'-'+a,text in md)
review=rows(R/'outputs/review/error_recheck_v2_20260914/rechecked_8.jsonl');counts=Counter(x['complete']for x in review)
check('8-case review counts',len(review)==8 and counts=={1:3,0:5})
for cid,p,n in [('CTX-001',2,2),('CTX-002',1,3),('CTX-003',0,3)]:
 rr=[x for x in review if x['case_id']==cid];check('additional case '+cid,len(rr)==n and sum(x['complete']for x in rr)==p)
check('AMF distortion and config coverage',[x['distorted']for x in review if x['case_id']=='CTX-003']==[True]*3 and all(x['requirements']['CONFIGURED_VERIFICATION']==0 for x in review if x['case_id']=='CTX-003'))
check('NSSAI two eCall incomplete',sum(x['requirements']['ECALL_EXCEPTION']==0 for x in review if x['case_id']=='CTX-002')==2)
raw=rows(R/'experiments/runs/error_replication_remaining8_20260914/parsed_responses.jsonl');last=next(x for x in raw if x['slot_id']=='ERR-CTX-002-R3')['parsed']
check('NSSAI metadata contains temporary storage','임시 저장'in last['action'] and '임시 저장'in last['exception'])
check('limitations explicitly retained',all(x in md for x in ['신뢰구간이 아니다','사후','동일 AI','독립·지연 검증이 아니다','주 분석과 합산하지 않았다']))
with pdfplumber.open(B/(stem+'.pdf'))as pdf:
 check('A4 two-page PDF',len(pdf.pages)==2 and all(abs(p.width-595.3)<1 and abs(p.height-841.9)<1 for p in pdf.pages))
 text='\n'.join(p.crop((x0,488 if i==0 else 0,x1,p.height)).extract_text() for i,p in enumerate(pdf.pages) for x0,x1 in [(0,p.width/2),(p.width/2,p.width)]);norm=lambda s:re.sub(r'\s+','',s)
 for p in d.paragraphs:
  if p.text.strip():check('DOCX paragraph present in PDF: '+p.text[:16],norm(p.text)in norm(text))
 check('selectable Korean and no replacement glyph','규범'in text and '\ufffd'not in text)
figure=d.inline_shapes[0];rid=figure._inline.xpath('.//a:blip')[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
check('embedded figure matches preserved plot',hashlib.sha256(d.part.related_parts[rid].blob).hexdigest()==h(B/'cohort_bounds_paper_v5.png'))
report={'status':'CONSISTENCY_CHECK_PASS_WITH_STATED_LIMITATIONS','checked_at':datetime.datetime.now(datetime.timezone.utc).isoformat(),'checks':checks,'blocking_discrepancies':[],'source_or_paper_edits':0,'new_generation_calls':0,'independent_human_validation':False,'paper_sha256':{ext:h(B/(stem+ext))for ext in ['.pdf','.docx','.md']},'scope':'Internal consistency and provenance, not semantic truth proof, acceptance prediction, or renewed full literature review.'}
(O/'checks.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n');print(json.dumps({'status':report['status'],'checks':len(checks),'blocking_discrepancies':0,'paper_unchanged':True}))
